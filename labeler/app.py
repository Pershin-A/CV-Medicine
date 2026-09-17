from pathlib import Path
from datetime import datetime
import hashlib
import re
import os
import shutil
import tempfile
import warnings

import numpy as np
import pandas as pd
import pydicom
import streamlit as st
from openpyxl import load_workbook
from PIL import Image
from pydicom.pixel_data_handlers.util import apply_voi_lut


APP_VERSION = "4.5-fast-clinical"

warnings.filterwarnings(
    "ignore",
    message=r"Invalid value for VR UI:.*",
    category=UserWarning,
    module=r"pydicom\.valuerep",
)
warnings.filterwarnings(
    "ignore",
    message=r"The value length .* exceeds the maximum length .* allowed for VR LO.*",
    category=UserWarning,
    module=r"pydicom\.valuerep",
)

DATA_ROOT = Path(os.getenv("DATA_ROOT", "/data"))
OUTPUT_ROOT = Path(os.getenv("OUTPUT_ROOT", "/output"))
REFERENCE_XLSX = Path(os.getenv("REFERENCE_XLSX", "/reference/разметка.xlsx"))

EXPECTED_DICOM_COUNT = int(os.getenv("EXPECTED_DICOM_COUNT", "0") or 0)
EXPECTED_DATASET_FINGERPRINT = os.getenv(
    "EXPECTED_DATASET_FINGERPRINT", ""
).strip().lower()
STRICT_DATASET_CHECK = os.getenv(
    "STRICT_DATASET_CHECK", "1"
).strip().lower() not in {"0", "false", "no"}

PRIVATE_GROUP = 0x0011
PRIVATE_CREATOR = "DXA_MANUAL_LABELER"

OFFSET_LABEL = 0x01
OFFSET_NOTES = 0x02
OFFSET_ANNOTATOR = 0x03
OFFSET_DATETIME = 0x04
OFFSET_SIDE = 0x05
OFFSET_METAL = 0x06
OFFSET_FRACTURE = 0x07
OFFSET_SPINE_ISSUE = 0x08

METAL_OPTIONS = {"Не размечено": "", "Нет металла": "0", "Есть металл": "1"}
METAL_OPTIONS_REVERSE = {v: k for k, v in METAL_OPTIONS.items()}

FRACTURE_OPTIONS = {"Не размечено": "", "Нет перелома": "0", "Есть перелом": "1"}
FRACTURE_OPTIONS_REVERSE = {v: k for k, v in FRACTURE_OPTIONS.items()}

SPINE_ISSUE_OPTIONS = {
    "Не размечено": "",
    "Нет": "NONE",
    "Сколиоз": "SCOLIOSIS",
    "Люмбализация": "LUMBARIZATION",
    "Сколиоз + люмбализация": "BOTH",
}
SPINE_ISSUE_OPTIONS_REVERSE = {v: k for k, v in SPINE_ISSUE_OPTIONS.items()}

CLASS_OPTIONS = {
    "Не определено": ("UNKNOWN", ""),
    "Позвоночник": ("SPINE", ""),
    "Нога — левая": ("LEG", "LEFT"),
    "Нога — правая": ("LEG", "RIGHT"),
    "Нога — сторона не указана": ("LEG", ""),
}


def class_option_from_annotation(label, side):
    label = str(label or "UNKNOWN").strip().upper()
    side = str(side or "").strip().upper()
    if label == "SPINE":
        return "Позвоночник"
    if label == "LEG" and side == "LEFT":
        return "Нога — левая"
    if label == "LEG" and side == "RIGHT":
        return "Нога — правая"
    if label == "LEG":
        return "Нога — сторона не указана"
    return "Не определено"


st.set_page_config(
    page_title="DXA DICOM Labeler",
    page_icon="🩻",
    layout="wide",
)


# ============================================================
# Пути, список файлов, fingerprint
# ============================================================

def normalize_relpath(value) -> str:
    return str(value).replace("\\", "/").lstrip("./")


def find_dicoms(root: Path):
    return sorted(
        (p for p in root.rglob("*") if p.is_file() and p.suffix.lower() == ".dcm"),
        key=lambda p: normalize_relpath(p.relative_to(root)).casefold(),
    )


@st.cache_data(show_spinner=False)
def cached_dataset_index(root_str: str):
    """Сканируем bind mount /data только один раз до ручного refresh."""
    root = Path(root_str)
    paths = find_dicoms(root)
    relpaths, sizes = [], []
    for path in paths:
        relpaths.append(normalize_relpath(path.relative_to(root)))
        try:
            sizes.append(path.stat().st_size)
        except OSError:
            sizes.append(-1)
    return relpaths, sizes


def relative_path_string(source_path: Path) -> str:
    return normalize_relpath(source_path.relative_to(DATA_ROOT))


def dataset_manifest(files):
    rows = []
    for path in files:
        try:
            size = path.stat().st_size
        except OSError:
            size = -1

        rows.append({
            "relative_path": relative_path_string(path),
            "size_bytes": size,
        })

    return pd.DataFrame(rows)


def dataset_fingerprint(manifest_df: pd.DataFrame) -> str:
    h = hashlib.sha256()

    for row in manifest_df.itertuples(index=False):
        line = f"{row.relative_path}|{row.size_bytes}\n"
        h.update(line.encode("utf-8", errors="replace"))

    return h.hexdigest()


def output_path_for(source_path: Path) -> Path:
    return OUTPUT_ROOT / source_path.relative_to(DATA_ROOT)


def effective_path_for(source_path: Path) -> Path:
    annotated = output_path_for(source_path)
    return annotated if annotated.exists() else source_path


# ============================================================
# CSV разметки: labels.csv = основной реестр
# ============================================================

def labels_csv_path() -> Path:
    return OUTPUT_ROOT / "labels.csv"


def history_csv_path() -> Path:
    return OUTPUT_ROOT / "labels_history.csv"


def read_labels_df() -> pd.DataFrame:
    path = labels_csv_path()

    if not path.exists():
        return pd.DataFrame(columns=[
            "relative_path",
            "label",
            "notes",
            "annotator",
            "annotated_at",
            "output_path",
            "side",
            "metal",
            "fracture",
            "spine_issue",
        ])

    try:
        df = pd.read_csv(path, dtype=str).fillna("")
    except Exception as e:
        st.error(
            f"Не удалось прочитать существующий labels.csv: {e}\n\n"
            "Файл не перезаписывается автоматически. "
            "Проверьте labels.csv и labels.csv.bak в папке результатов."
        )
        st.stop()

    if "relative_path" not in df.columns:
        st.error(
            "В существующем labels.csv нет колонки relative_path. "
            "Чтобы не повредить уже сделанную разметку, сохранение остановлено."
        )
        st.stop()

    df["relative_path"] = df["relative_path"].map(normalize_relpath)
    if "side" not in df.columns:
        df["side"] = ""
    df["side"] = df["side"].astype(str).str.strip().str.upper()
    if "metal" not in df.columns:
        df["metal"] = ""
    df["metal"] = df["metal"].astype(str).str.strip()
    if "fracture" not in df.columns:
        df["fracture"] = ""
    df["fracture"] = df["fracture"].astype(str).str.strip()
    if "spine_issue" not in df.columns:
        df["spine_issue"] = ""
    df["spine_issue"] = df["spine_issue"].astype(str).str.strip().str.upper()

    # Если по одному пути по какой-то причине несколько строк,
    # последняя строка считается текущим состоянием.
    df = df.drop_duplicates(
        subset=["relative_path"],
        keep="last",
    )

    return df.reset_index(drop=True)


def labels_lookup(df: pd.DataFrame):
    if df.empty:
        return {}

    return {
        normalize_relpath(row["relative_path"]): row
        for _, row in df.iterrows()
    }


def atomic_write_csv(df: pd.DataFrame, destination: Path):
    destination.parent.mkdir(parents=True, exist_ok=True)

    fd, tmp_name = tempfile.mkstemp(
        prefix=destination.stem + "_",
        suffix=".tmp",
        dir=str(destination.parent),
    )
    os.close(fd)

    tmp_path = Path(tmp_name)

    try:
        df.to_csv(
            tmp_path,
            index=False,
            encoding="utf-8-sig",
        )
        os.replace(tmp_path, destination)
    finally:
        if tmp_path.exists():
            tmp_path.unlink(missing_ok=True)


def append_history_row(row: dict):
    path = history_csv_path()
    path.parent.mkdir(parents=True, exist_ok=True)

    new_df = pd.DataFrame([row])

    if path.exists():
        try:
            old = pd.read_csv(path, dtype=str).fillna("")
            out = pd.concat([old, new_df], ignore_index=True)
        except Exception:
            # Не трогаем повреждённую историю; создаём отдельный recovery-файл.
            recovery = OUTPUT_ROOT / (
                "labels_history_recovery_"
                + datetime.now().strftime("%Y%m%d_%H%M%S")
                + ".csv"
            )
            new_df.to_csv(
                recovery,
                index=False,
                encoding="utf-8-sig",
            )
            return
    else:
        out = new_df

    atomic_write_csv(out, path)


def update_labels_csv(
    relative_path,
    label,
    notes,
    annotator,
    annotated_at,
    output_path,
    side="",
    metal="",
    fracture="",
    spine_issue="",
):
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)

    csv_path = labels_csv_path()
    backup_path = OUTPUT_ROOT / "labels.csv.bak"

    df = read_labels_df()

    # Перед каждым обновлением сохраняем предыдущую рабочую версию.
    if csv_path.exists():
        shutil.copy2(csv_path, backup_path)

    relative_path = normalize_relpath(relative_path)

    new_row = {
        "relative_path": relative_path,
        "label": label,
        "notes": notes,
        "annotator": annotator,
        "annotated_at": annotated_at,
        "output_path": normalize_relpath(
            Path(output_path).relative_to(OUTPUT_ROOT)
        ),
        "side": str(side or "").strip().upper(),
        "metal": str(metal or "").strip(),
        "fracture": str(fracture or "").strip(),
        "spine_issue": str(spine_issue or "").strip().upper(),
    }

    if not df.empty:
        df = df[
            df["relative_path"].map(normalize_relpath)
            != relative_path
        ]

    df = pd.concat(
        [df, pd.DataFrame([new_row])],
        ignore_index=True,
    )

    df["relative_path"] = df["relative_path"].map(normalize_relpath)
    df = df.sort_values(
        "relative_path",
        key=lambda s: s.str.casefold(),
    )

    atomic_write_csv(df, csv_path)
    append_history_row(new_row)


# ============================================================
# Private DICOM tags
# ============================================================

def get_private_block(ds, create=False):
    try:
        return ds.private_block(
            PRIVATE_GROUP,
            PRIVATE_CREATOR,
            create=create,
        )
    except Exception:
        return None


def get_private_value(ds, offset, default=""):
    block = get_private_block(ds, create=False)

    if block is None:
        return default

    try:
        tag = block.get_tag(offset)
    except Exception:
        return default

    if tag in ds:
        value = ds[tag].value
        if value is None:
            return default
        if isinstance(value, bytes):
            value = value.decode("utf-8", errors="replace")
        return str(value).replace("\x00", "").strip()

    return default


def annotation_from_dicom(source_path: Path):
    path = effective_path_for(source_path)

    try:
        ds = pydicom.dcmread(
            path,
            stop_before_pixels=True,
            force=True,
        )
    except Exception:
        return {
            "label": "UNKNOWN",
            "notes": "",
            "annotator": "",
            "datetime": "",
            "side": "",
            "metal": "",
            "fracture": "",
            "spine_issue": "",
        }

    return {
        "label": get_private_value(
            ds, OFFSET_LABEL, "UNKNOWN"
        ),
        "notes": get_private_value(
            ds, OFFSET_NOTES, ""
        ),
        "annotator": get_private_value(
            ds, OFFSET_ANNOTATOR, ""
        ),
        "datetime": get_private_value(
            ds, OFFSET_DATETIME, ""
        ),
        "side": get_private_value(ds, OFFSET_SIDE, "").upper(),
        "metal": get_private_value(ds, OFFSET_METAL, ""),
        "fracture": get_private_value(ds, OFFSET_FRACTURE, ""),
        "spine_issue": get_private_value(ds, OFFSET_SPINE_ISSUE, "").upper(),
    }


def read_existing_annotation(
    source_path: Path,
    lookup: dict,
):
    """
    labels.csv считается главным источником состояния.
    Это важно для уже размеченных данных:
    даже если размеченная DICOM-копия отсутствует/перемещена,
    существующая строка labels.csv не теряется.
    """
    rel = relative_path_string(source_path)

    if rel in lookup:
        row = lookup[rel]

        return {
            "label": str(row.get("label", "UNKNOWN")) or "UNKNOWN",
            "notes": str(row.get("notes", "")),
            "annotator": str(row.get("annotator", "")),
            "datetime": str(row.get("annotated_at", "")),
            "side": str(row.get("side", "")).strip().upper(),
            "metal": str(row.get("metal", "")).strip(),
            "fracture": str(row.get("fracture", "")).strip(),
            "spine_issue": str(row.get("spine_issue", "")).strip().upper(),
            "source": "labels.csv",
        }

    result = annotation_from_dicom(source_path)
    result["source"] = "DICOM"
    return result


# ============================================================
# DICOM image / metadata
# ============================================================

def dicom_to_image(ds):
    arr = ds.pixel_array.astype(np.float32)

    slope = float(getattr(ds, "RescaleSlope", 1.0))
    intercept = float(getattr(ds, "RescaleIntercept", 0.0))
    arr = arr * slope + intercept

    try:
        arr = apply_voi_lut(arr, ds).astype(np.float32)
    except Exception:
        pass

    finite = np.isfinite(arr)

    if not finite.any():
        raise ValueError(
            "Изображение не содержит конечных значений."
        )

    lo = float(np.nanpercentile(arr[finite], 1))
    hi = float(np.nanpercentile(arr[finite], 99))

    if hi <= lo:
        lo = float(np.nanmin(arr[finite]))
        hi = float(np.nanmax(arr[finite]))

    if hi <= lo:
        image = np.zeros(arr.shape, dtype=np.uint8)
    else:
        image = np.clip(
            (arr - lo) / (hi - lo),
            0,
            1,
        )
        image = (image * 255).astype(np.uint8)

    if (
        str(getattr(ds, "PhotometricInterpretation", ""))
        == "MONOCHROME1"
    ):
        image = 255 - image

    return Image.fromarray(image)


@st.cache_data(show_spinner=False, max_entries=16)
def cached_preview(path_str: str, mtime_ns: int):
    ds = pydicom.dcmread(path_str, force=True)
    return np.asarray(dicom_to_image(ds))


def safe_value(element):
    try:
        value = element.value
    except Exception as e:
        return f"<error: {e}>"

    if isinstance(value, bytes):
        return f"<binary: {len(value)} bytes>"

    text = str(value)

    if len(text) > 1000:
        return (
            text[:1000]
            + f" ... <truncated; length={len(text)}>"
        )

    return text


def metadata_dataframe(ds):
    rows = []

    if getattr(ds, "file_meta", None):
        for element in ds.file_meta:
            rows.append({
                "Section": "FileMeta",
                "Tag": str(element.tag),
                "Keyword": element.keyword or "",
                "Name": element.name,
                "VR": element.VR,
                "Value": safe_value(element),
            })

    for element in ds.iterall():
        if element.tag == (0x7FE0, 0x0010):
            value = "<Pixel Data omitted>"
        else:
            value = safe_value(element)

        rows.append({
            "Section": "Dataset",
            "Tag": str(element.tag),
            "Keyword": element.keyword or "",
            "Name": element.name,
            "VR": element.VR,
            "Value": value,
        })

    return pd.DataFrame(rows)


# ============================================================
# Безопасное сохранение DICOM
# ============================================================

def atomic_save_dicom(ds, out_path: Path):
    out_path.parent.mkdir(parents=True, exist_ok=True)

    temp_path = out_path.with_name(
        out_path.name + ".tmp"
    )

    try:
        ds.save_as(
            temp_path,
            enforce_file_format=True,
        )
        os.replace(temp_path, out_path)
    finally:
        temp_path.unlink(missing_ok=True)


def save_annotation(
    source_path: Path,
    label: str,
    notes: str,
    annotator: str,
    side: str = "",
    metal: str = "",
    fracture: str = "",
    spine_issue: str = "",
):
    # labels.csv — главный реестр. При сохранении всегда создаём копию
    # из чистого исходного DICOM, чтобы не накапливать старые private tags.
    ds = pydicom.dcmread(
        source_path,
        force=True,
    )

    block = get_private_block(
        ds,
        create=True,
    )

    if block is None:
        raise RuntimeError(
            "Не удалось создать private DICOM block "
            "для разметки."
        )

    label_tag = block.get_tag(OFFSET_LABEL)
    notes_tag = block.get_tag(OFFSET_NOTES)
    annotator_tag = block.get_tag(OFFSET_ANNOTATOR)
    datetime_tag = block.get_tag(OFFSET_DATETIME)
    side_tag = block.get_tag(OFFSET_SIDE)
    metal_tag = block.get_tag(OFFSET_METAL)
    fracture_tag = block.get_tag(OFFSET_FRACTURE)
    spine_issue_tag = block.get_tag(OFFSET_SPINE_ISSUE)

    now = datetime.now()
    dicom_dt = now.strftime("%Y%m%d%H%M%S.%f")
    iso_dt = now.isoformat(timespec="seconds")

    ds.add_new(
        label_tag,
        "CS",
        label,
    )
    ds.add_new(
        notes_tag,
        "LT",
        notes or "",
    )
    ds.add_new(
        annotator_tag,
        "LO",
        annotator or "",
    )
    ds.add_new(datetime_tag, "DT", dicom_dt)
    ds.add_new(side_tag, "CS", side if label == "LEG" else "")
    ds.add_new(metal_tag, "CS", metal if label == "LEG" else "")
    ds.add_new(fracture_tag, "CS", fracture if label == "LEG" else "")
    ds.add_new(spine_issue_tag, "CS", spine_issue if label == "SPINE" else "")

    out_path = output_path_for(source_path)

    # 1. Сначала безопасно пишем размеченную DICOM-копию.
    atomic_save_dicom(ds, out_path)

    # 2. Затем atomically обновляем labels.csv + backup + history.
    relative_path = relative_path_string(source_path)

    update_labels_csv(
        relative_path=relative_path,
        label=label,
        notes=notes or "",
        annotator=annotator or "",
        annotated_at=iso_dt,
        output_path=out_path,
        side=side if label == "LEG" else "",
        metal=metal if label == "LEG" else "",
        fracture=fracture if label == "LEG" else "",
        spine_issue=spine_issue if label == "SPINE" else "",
    )

    return out_path



# ============================================================
# Справочная разметка из разметка.xlsx
# ============================================================

def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and np.isnan(value):
        return ""
    return str(value).strip()


def _norm_ru(value):
    return " ".join(_cell_text(value).lower().replace("ё", "е").replace("\xa0", " ").split())


def _looks_like_study_uid(value):
    text = _cell_text(value)
    return bool(re.fullmatch(r"(?:2\.25\.\d+|1\.2(?:\.\d+)+)", text))


def _is_one(value):
    text = _cell_text(value).replace(",", ".")
    try:
        return float(text) == 1.0
    except Exception:
        return text == "1"


def _header_area(parts):
    texts = [_norm_ru(x) for x in parts if _cell_text(x)]
    joined = " | ".join(texts)
    if ("прав" in joined and "бедр" in joined) or "правого бедра" in joined:
        return "RIGHT_LEG"
    if ("лев" in joined and "бедр" in joined) or "левого бедра" in joined:
        return "LEFT_LEG"
    if "позвоноч" in joined:
        return "SPINE"
    return None


def _is_generic_header(text):
    n = _norm_ru(text)
    generic = {
        "", "№", "номер", "study", "итог", "общий", "позвоночник",
        "проксимальный отдел правого бедра", "проксимальный отдел левого бедра",
    }
    return n in generic


@st.cache_data(show_spinner=False)
def load_reference_excel(path_str: str, mtime_ns: int):
    """Возвращает простой dict study -> область/единицы/итог/комментарий."""
    path = Path(path_str)
    if not path.exists() or not path.is_file():
        return {}, {"status": "missing", "path": str(path)}

    try:
        wb = load_workbook(path, data_only=True, read_only=False)
        ws = wb.active
    except Exception as e:
        return {}, {"status": "error", "path": str(path), "error": str(e)}

    max_row, max_col = ws.max_row, ws.max_column
    grid = [[ws.cell(r, c).value for c in range(1, max_col + 1)] for r in range(1, max_row + 1)]

    # Восстанавливаем значения merged cells во всех ячейках диапазона.
    for merged in ws.merged_cells.ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for r in range(merged.min_row, merged.max_row + 1):
            for c in range(merged.min_col, merged.max_col + 1):
                grid[r - 1][c - 1] = value

    uid_counts = []
    for c in range(max_col):
        uid_counts.append(sum(_looks_like_study_uid(grid[r][c]) for r in range(max_row)))
    study_col = int(np.argmax(uid_counts)) if uid_counts else 0
    if not uid_counts or uid_counts[study_col] == 0:
        return {}, {"status": "error", "path": str(path), "error": "Не найден столбец study UID"}

    data_rows = [r for r in range(max_row) if _looks_like_study_uid(grid[r][study_col])]
    if not data_rows:
        return {}, {"status": "error", "path": str(path), "error": "Не найдены строки данных"}
    first_data = min(data_rows)
    header_rows = grid[:first_data]

    descriptors = []
    for c in range(max_col):
        parts = []
        for r in range(first_data):
            text = _cell_text(header_rows[r][c])
            if text and (not parts or parts[-1] != text):
                parts.append(text)
        norms = [_norm_ru(x) for x in parts]
        joined = " | ".join(norms)
        area = _header_area(parts)
        is_comment = "комментар" in joined
        is_total = any(n == "итог" or n.startswith("итог ") for n in norms)
        is_general = any(n == "общий" or n.startswith("общий ") for n in norms)

        leaf = ""
        for part in reversed(parts):
            n = _norm_ru(part)
            if not _is_generic_header(part) and "комментар" not in n:
                leaf = part
                break

        descriptors.append({
            "col": c, "parts": parts, "area": area, "is_comment": is_comment,
            "is_total": is_total, "is_general": is_general, "leaf": leaf,
        })

    result = {}
    for r in data_rows:
        study = _cell_text(grid[r][study_col])
        record = {
            "SPINE": {"ones": [], "total": ""},
            "RIGHT_LEG": {"ones": [], "total": ""},
            "LEFT_LEG": {"ones": [], "total": ""},
            "comment": "",
        }
        comments = []
        for d in descriptors:
            value = grid[r][d["col"]]
            text = _cell_text(value)
            if d["is_comment"]:
                if text:
                    comments.append(text)
                continue
            area = d["area"]
            if area not in record:
                continue
            if d["is_total"]:
                if text and not record[area]["total"]:
                    record[area]["total"] = text
                continue
            if d["is_general"]:
                continue
            if d["leaf"] and _is_one(value):
                if d["leaf"] not in record[area]["ones"]:
                    record[area]["ones"].append(d["leaf"])

        record["comment"] = " | ".join(dict.fromkeys(comments))
        result[study] = record

    diagnostics = {
        "status": "ok", "path": str(path), "studies": len(result),
        "study_column_index": study_col + 1,
        "first_data_row": first_data + 1,
    }
    return result, diagnostics


def reference_area_for(label, side):
    if label == "SPINE":
        return "SPINE", "Позвоночник"
    if label == "LEG" and side == "RIGHT":
        return "RIGHT_LEG", "Проксимальный отдел правого бедра"
    if label == "LEG" and side == "LEFT":
        return "LEFT_LEG", "Проксимальный отдел левого бедра"
    return None, None


def render_reference_info(reference_record, label, side, study_id):
    st.markdown("#### Информация из `разметка.xlsx`")
    st.caption(f"study: `{study_id}`")

    if reference_record is None:
        st.info("Для этого study строка в разметка.xlsx не найдена.")
        st.write("**Комментарий:** —")
        return

    area_key, area_title = reference_area_for(label, side)
    if area_key is None:
        st.info("Выберите позвоночник или сторону ноги, чтобы показать поля для соответствующей области.")
    else:
        area = reference_record.get(area_key, {"ones": [], "total": ""})
        st.write(f"**Область:** {area_title}")
        if area.get("ones"):
            st.write("**Поля со значением 1:**")
            for name in area["ones"]:
                st.write(f"• {name}")
        else:
            st.write("**Поля со значением 1:** нет")
        total = area.get("total", "")
        st.write(f"**Итог для области:** **{total if total != '' else '—'}**")

    comment = reference_record.get("comment", "")
    st.write(f"**Комментарий:** {comment if comment else '—'}")


# ============================================================
# Навигация
# ============================================================

def move_to(index: int, files):
    index = max(
        0,
        min(index, len(files) - 1),
    )

    st.session_state.current_idx = index

    # Значение selectbox изменяем на следующем rerun,
    # до создания самого widget.
    st.session_state.pending_file_select = (
        relative_path_string(files[index])
    )


# ============================================================
# Startup / dataset validation
# ============================================================

st.title(f"🩻 DXA DICOM manual labeler — {APP_VERSION}")
st.caption(
    "Исходные DICOM читаются только из /data. "
    "Разметка хранится в /output. "
    "labels.csv является основным реестром разметки. "
    f"Для ног размечаются сторона, металл и перелом; для позвоночника — тип проблемы. Версия: {APP_VERSION}."
)

if not DATA_ROOT.exists():
    st.error(
        f"Каталог с данными не найден: {DATA_ROOT}\n\n"
        "Проверьте DXA_DATA_DIR в файле .env."
    )
    st.stop()

OUTPUT_ROOT.mkdir(
    parents=True,
    exist_ok=True,
)

if REFERENCE_XLSX.exists() and REFERENCE_XLSX.is_file():
    reference_by_study, reference_diag = load_reference_excel(
        str(REFERENCE_XLSX), REFERENCE_XLSX.stat().st_mtime_ns
    )
else:
    reference_by_study, reference_diag = {}, {
        "status": "missing", "path": str(REFERENCE_XLSX)
    }

relative_options, file_sizes = cached_dataset_index(str(DATA_ROOT))
files = [DATA_ROOT / Path(*Path(rel).parts) for rel in relative_options]
manifest_df = pd.DataFrame({"relative_path": relative_options, "size_bytes": file_sizes})
fingerprint = dataset_fingerprint(manifest_df)

if not files:
    st.error(f"В {DATA_ROOT} не найдено файлов *.dcm")
    st.stop()

labels_df = read_labels_df()
lookup = labels_lookup(labels_df)

current_paths = set(relative_options)

annotated_current = (
    set(labels_df["relative_path"].map(normalize_relpath))
    & current_paths
    if not labels_df.empty
    else set()
)

n_annotated = len(annotated_current)

if "side" in labels_df.columns:
    leg_rows = labels_df[labels_df["label"].astype(str).str.upper().eq("LEG")]
    n_leg_side = leg_rows["side"].astype(str).str.upper().isin(["LEFT", "RIGHT"]).sum()
    n_leg_total = len(leg_rows)
else:
    n_leg_side, n_leg_total = 0, 0

if "metal" in labels_df.columns:
    n_metal_labeled = leg_rows["metal"].astype(str).isin(["0", "1"]).sum() if n_leg_total else 0
else:
    n_metal_labeled = 0

if "fracture" in labels_df.columns:
    n_fracture_labeled = leg_rows["fracture"].astype(str).isin(["0", "1"]).sum() if n_leg_total else 0
else:
    n_fracture_labeled = 0

spine_rows = labels_df[labels_df["label"].astype(str).str.upper().eq("SPINE")]
n_spine_total = len(spine_rows)
if "spine_issue" in labels_df.columns:
    n_spine_issue_labeled = spine_rows["spine_issue"].astype(str).str.upper().isin(
        ["NONE", "SCOLIOSIS", "LUMBARIZATION", "BOTH"]
    ).sum()
else:
    n_spine_issue_labeled = 0

count_ok = (
    EXPECTED_DICOM_COUNT <= 0
    or len(files) == EXPECTED_DICOM_COUNT
)

fingerprint_ok = (
    not EXPECTED_DATASET_FINGERPRINT
    or fingerprint.lower()
    == EXPECTED_DATASET_FINGERPRINT
)

dataset_ok = count_ok and fingerprint_ok


# ============================================================
# Session state before widgets
# ============================================================

if "current_idx" not in st.session_state:
    st.session_state.current_idx = 0

st.session_state.current_idx = min(
    st.session_state.current_idx,
    len(files) - 1,
)

if "file_select" not in st.session_state:
    st.session_state.file_select = relative_options[
        st.session_state.current_idx
    ]

if "pending_file_select" in st.session_state:
    pending_value = st.session_state.pop(
        "pending_file_select"
    )

    if pending_value in relative_options:
        st.session_state.file_select = pending_value
        st.session_state.current_idx = (
            relative_options.index(pending_value)
        )

# Если список данных изменился и старое значение selectbox
# больше не существует, аккуратно возвращаемся к текущему индексу.
if st.session_state.file_select not in relative_options:
    st.session_state.current_idx = min(
        st.session_state.current_idx,
        len(relative_options) - 1,
    )
    st.session_state.file_select = relative_options[
        st.session_state.current_idx
    ]


def on_file_selected():
    selected = st.session_state.file_select

    try:
        st.session_state.current_idx = (
            relative_options.index(selected)
        )
    except ValueError:
        st.session_state.current_idx = 0


# ============================================================
# Sidebar
# ============================================================

with st.sidebar:
    st.header("Навигация")
    st.success(f"Версия интерфейса: {APP_VERSION}")
    if st.button("🔄 Пересканировать /data", width="stretch"):
        st.cache_data.clear()
        st.rerun()

    st.selectbox(
        "Выберите DICOM",
        options=relative_options,
        key="file_select",
        on_change=on_file_selected,
    )

    idx = st.session_state.current_idx

    st.write(
        f"Файл **{idx + 1} / {len(files)}**"
    )

    prev_col, next_col = st.columns(2)

    with prev_col:
        if st.button(
            "← Пред.",
            width="stretch",
            disabled=(idx == 0),
        ):
            move_to(
                idx - 1,
                files,
            )
            st.rerun()

    with next_col:
        if st.button(
            "След. →",
            width="stretch",
            disabled=(idx >= len(files) - 1),
        ):
            move_to(
                idx + 1,
                files,
            )
            st.rerun()

    st.divider()

    st.metric(
        "Размечено в текущем наборе",
        f"{n_annotated} / {len(files)}",
    )

    st.progress(min(n_annotated / len(files), 1.0))
    st.metric("Сторона размечена у ног", f"{n_leg_side} / {n_leg_total}")
    st.metric("Металл размечен у ног", f"{n_metal_labeled} / {n_leg_total}")
    st.metric("Перелом размечен у ног", f"{n_fracture_labeled} / {n_leg_total}")
    st.metric("Проблема позвоночника", f"{n_spine_issue_labeled} / {n_spine_total}")

    if len(labels_df) != n_annotated:
        st.caption(
            f"Всего строк в labels.csv: {len(labels_df)}. "
            "Некоторые строки могут относиться к файлам, "
            "которых нет в текущем /data."
        )

    st.divider()

    st.subheader("Проверка набора")

    if count_ok:
        st.success(
            f"DICOM в контейнере: {len(files)}"
        )
    else:
        st.error(
            f"DICOM в контейнере: {len(files)}, "
            f"ожидалось: {EXPECTED_DICOM_COUNT}"
        )

    st.code(
        fingerprint[:16],
        language=None,
    )
    st.caption(
        "Fingerprint должен совпадать у всех разметчиков, "
        "если набор файлов один и тот же."
    )

    st.download_button(
        "Скачать manifest.csv",
        data=manifest_df.to_csv(
            index=False
        ).encode("utf-8-sig"),
        file_name="dataset_manifest.csv",
        mime="text/csv",
        width="stretch",
    )

    with st.expander(
        "Технические пути"
    ):
        st.write(
            f"DATA_ROOT: `{DATA_ROOT}`"
        )
        st.write(
            f"OUTPUT_ROOT: `{OUTPUT_ROOT}`"
        )
        st.write(f"labels.csv: `{labels_csv_path()}`")
        st.write(f"reference xlsx: `{REFERENCE_XLSX}`")
        if reference_diag.get("status") == "ok":
            st.success(f"Excel: {reference_diag.get('studies', 0)} study")
        elif reference_diag.get("status") == "missing":
            st.warning("Excel не смонтирован")
        else:
            st.error(f"Excel error: {reference_diag.get('error', '')}")


# ============================================================
# Stop new labeling if dataset check failed
# ============================================================

if not dataset_ok:
    st.error(
        "Проверка набора данных не пройдена. "
        "Разметка остановлена, чтобы не продолжать работу "
        "по неполному или другому списку файлов.\n\n"
        "Существующие labels.csv и размеченные DICOM НЕ изменены."
    )

    if STRICT_DATASET_CHECK:
        st.stop()
    else:
        st.warning(
            "STRICT_DATASET_CHECK=0: работа разрешена "
            "несмотря на несовпадение."
        )


# ============================================================
# Current file
# ============================================================

current_idx = st.session_state.current_idx
source_path = files[current_idx]
relative_path = relative_path_string(source_path)
existing = read_existing_annotation(source_path, lookup)

top_left, top_right = st.columns(
    [1.05, 0.95],
    gap="large",
)

with top_left:
    st.subheader("Изображение")
    st.code(
        relative_path,
        language=None,
    )

    try:
        image = cached_preview(str(source_path), source_path.stat().st_mtime_ns)
        st.image(image, caption=relative_path, width="stretch")
    except Exception as e:
        st.error(
            "Не удалось декодировать PixelData.\n\n"
            f"{e}\n\n"
            "Если файл сжат, проверьте pylibjpeg."
        )


with top_right:
    st.subheader("Разметка")

    if relative_path in annotated_current:
        st.success(
            "Файл уже размечен. "
            f"Источник текущей метки: {existing['source']}."
        )
    elif output_path_for(source_path).exists():
        st.info(
            "Есть размеченная DICOM-копия, "
            "но записи для неё нет в текущем labels.csv."
        )

    existing_label = str(existing.get("label", "UNKNOWN") or "UNKNOWN").strip().upper()
    class_options = {"Не определено": "UNKNOWN", "Позвоночник": "SPINE", "Нога": "LEG"}
    initial_class = "Позвоночник" if existing_label == "SPINE" else ("Нога" if existing_label == "LEG" else "Не определено")

    widget_suffix = hashlib.sha1(relative_path.encode("utf-8", errors="replace")).hexdigest()[:12]
    class_ui = st.radio(
        "Класс",
        options=list(class_options.keys()),
        index=list(class_options.keys()).index(initial_class),
        key=f"class_{widget_suffix}",
        horizontal=True,
    )
    current_label = class_options[class_ui]

    side_options = {"Не размечено": "", "Левая": "LEFT", "Правая": "RIGHT"}
    side_reverse = {v: k for k, v in side_options.items()}
    initial_side = side_reverse.get(str(existing.get("side", "")).strip().upper(), "Не размечено")
    side_ui = st.radio(
        "Сторона ноги",
        options=list(side_options.keys()),
        index=list(side_options.keys()).index(initial_side),
        key=f"side_{widget_suffix}",
        horizontal=True,
        disabled=(current_label != "LEG"),
    )
    current_side = side_options[side_ui] if current_label == "LEG" else ""

    existing_metal = str(existing.get("metal", "")).strip()
    initial_metal = METAL_OPTIONS_REVERSE.get(existing_metal, "Не размечено")
    metal_ui = st.radio(
        "Металл / имплант",
        options=list(METAL_OPTIONS.keys()),
        index=list(METAL_OPTIONS.keys()).index(initial_metal),
        key=f"metal_{widget_suffix}",
        horizontal=True,
        disabled=(current_label != "LEG"),
    )
    current_metal = METAL_OPTIONS[metal_ui] if current_label == "LEG" else ""

    existing_fracture = str(existing.get("fracture", "")).strip()
    initial_fracture = FRACTURE_OPTIONS_REVERSE.get(existing_fracture, "Не размечено")
    fracture_ui = st.radio(
        "Перелом",
        options=list(FRACTURE_OPTIONS.keys()),
        index=list(FRACTURE_OPTIONS.keys()).index(initial_fracture),
        key=f"fracture_{widget_suffix}",
        horizontal=True,
        disabled=(current_label != "LEG"),
    )
    current_fracture = FRACTURE_OPTIONS[fracture_ui] if current_label == "LEG" else ""

    existing_spine_issue = str(existing.get("spine_issue", "")).strip().upper()
    initial_spine_issue = SPINE_ISSUE_OPTIONS_REVERSE.get(existing_spine_issue, "Не размечено")
    spine_issue_ui = st.radio(
        "Проблема позвоночника",
        options=list(SPINE_ISSUE_OPTIONS.keys()),
        index=list(SPINE_ISSUE_OPTIONS.keys()).index(initial_spine_issue),
        key=f"spine_issue_{widget_suffix}",
        horizontal=False,
        disabled=(current_label != "SPINE"),
    )
    current_spine_issue = SPINE_ISSUE_OPTIONS[spine_issue_ui] if current_label == "SPINE" else ""

    study_id = relative_path.split("/", 1)[0]
    reference_record = reference_by_study.get(study_id)
    with st.container(border=True):
        render_reference_info(reference_record, current_label, current_side, study_id)

    if current_label == "LEG" and not current_side:
        st.warning("Для ноги выберите сторону: левая или правая.")

    annotator = st.text_input(
        "Разметчик",
        value=existing["annotator"],
        key=f"annotator_{widget_suffix}",
        placeholder="Например: Андрей",
    )

    notes = st.text_area(
        "Комментарий",
        value=existing["notes"],
        key=f"notes_{widget_suffix}",
        placeholder="Необязательно",
        height=100,
    )

    st.write(
        "Будет записано в private DICOM metadata:"
    )

    st.code(
        "\n".join([
            f"Private Creator: {PRIVATE_CREATOR}",
            f"ManualLabel: {current_label}",
            f"Side: {current_side or '<empty>'}",
            f"Metal: {current_metal or '<not labeled>'}",
            f"Fracture: {current_fracture or '<not labeled>'}",
            f"SpineIssue: {current_spine_issue or '<not labeled>'}",
            f"Notes: {notes or '<empty>'}",
            f"Annotator: {annotator or '<empty>'}",
        ]),
        language=None,
    )

    save_col, save_next_col = st.columns(2)

    with save_col:
        if st.button(
            "💾 Сохранить",
            type="primary",
            width="stretch",
        ):
            try:
                out_path = save_annotation(
                    source_path=source_path,
                    label=current_label,
                    notes=notes,
                    annotator=annotator,
                    side=current_side,
                    metal=current_metal,
                    fracture=current_fracture,
                    spine_issue=current_spine_issue,
                )
                st.success(
                    f"Сохранено:\n{out_path}"
                )
                st.rerun()
            except Exception as e:
                st.exception(e)

    with save_next_col:
        if st.button(
            "💾 Сохранить и следующий →",
            width="stretch",
        ):
            try:
                save_annotation(
                    source_path=source_path,
                    label=current_label,
                    notes=notes,
                    annotator=annotator,
                    side=current_side,
                    metal=current_metal,
                    fracture=current_fracture,
                    spine_issue=current_spine_issue,
                )

                if current_idx < len(files) - 1:
                    move_to(
                        current_idx + 1,
                        files,
                    )
                    st.rerun()
                else:
                    st.success(
                        "Последний файл сохранён."
                    )
            except Exception as e:
                st.exception(e)


# ============================================================
# Metadata table — читается только по запросу
# ============================================================

st.divider()
show_metadata = st.checkbox("Показать DICOM metadata", value=False)
if show_metadata:
    try:
        meta_ds = pydicom.dcmread(source_path, stop_before_pixels=True, force=True)
        metadata_df = metadata_dataframe(meta_ds)
        metadata_query = st.text_input("Поиск по metadata", key=f"metadata_search_{widget_suffix}")
        if metadata_query:
            q = metadata_query.casefold()
            mask = metadata_df.astype(str).apply(
                lambda col: col.str.casefold().str.contains(q, regex=False)
            ).any(axis=1)
            metadata_df = metadata_df[mask]
        st.dataframe(metadata_df, width="stretch", hide_index=True)
    except Exception as e:
        st.error(f"Не удалось прочитать metadata: {e}")
